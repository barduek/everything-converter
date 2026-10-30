"""Pruebas de src/converter/converters/docx/fonts.py: fuentes referenciadas, incrustadas y
reemplazadas de un DOCX (Req 7.4).

Los paquetes se arman a mano con las partes de WordprocessingML que importan en cada caso, y
uno se genera con python-docx para cubrir un documento real. Las pruebas que ejecutan fc-list
y fc-match se omiten fuera del contenedor.
"""

from __future__ import annotations

import html
import io
import zipfile
from collections.abc import Iterator, Mapping
from contextlib import contextmanager
from pathlib import Path

import docx
import pytest

from converter.converters.docx.fonts import (
    MAX_FONT_NAME_LENGTH,
    DocumentFonts,
    FontAnalysisError,
    fc_match_pattern,
    find_replacements,
    installed_families,
    match_font,
    missing_fonts,
    parse_installed_families,
    read_document_fonts,
)

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_A = "http://schemas.openxmlformats.org/drawingml/2006/main"
_W_STRICT = "http://purl.oclc.org/ooxml/wordprocessingml/main"
_R_STRICT = "http://purl.oclc.org/ooxml/officeDocument/relationships"
_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"


def _rels(*relationships: tuple[str, str, str], namespace: str = _R, external: str = "") -> str:
    """Parte de relaciones con (Id, tipo sin prefijo, destino); `external` es un Id externo."""
    items = "".join(
        f'<Relationship Id="{rel_id}" Type="{namespace}/{kind}" Target="{target}"'
        + (' TargetMode="External"' if rel_id == external else "")
        + "/>"
        for rel_id, kind, target in relationships
    )
    return f'<?xml version="1.0"?><Relationships xmlns="{_PKG}">{items}</Relationships>'


def _wml(root: str, body: str, *, w: str = _W, r: str = _R) -> str:
    return f'<w:{root} xmlns:w="{w}" xmlns:r="{r}">{body}</w:{root}>'


def _run(attributes: str) -> str:
    return f"<w:p><w:r><w:rPr><w:rFonts {attributes}/></w:rPr><w:t>hola</w:t></w:r></w:p>"


def _theme(*, major: str, minor: str, minor_ea: str = "") -> str:
    return (
        f'<a:theme xmlns:a="{_A}" name="Tema"><a:themeElements><a:fontScheme name="Esquema">'
        f'<a:majorFont><a:latin typeface="{major}"/><a:ea typeface=""/><a:cs typeface=""/>'
        f'<a:font script="Jpan" typeface="Fuente Jpan"/></a:majorFont>'
        f'<a:minorFont><a:latin typeface="{minor}"/><a:ea typeface="{minor_ea}"/>'
        f'<a:cs typeface=""/></a:minorFont>'
        "</a:fontScheme></a:themeElements>"
        # Las propiedades de texto del tema también usan a:latin, fuera del esquema de fuentes
        '<a:objectDefaults><a:defRPr><a:latin typeface="Fuente Objeto"/></a:defRPr>'
        "</a:objectDefaults></a:theme>"
    )


def _document(body: str) -> str:
    return _wml("document", f"<w:body>{body}</w:body>")


@contextmanager
def _package(parts: Mapping[str, str | bytes]) -> Iterator[zipfile.ZipFile]:
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w", zipfile.ZIP_DEFLATED) as archive:
        for name, data in parts.items():
            archive.writestr(name, data)
    with zipfile.ZipFile(buffer) as archive:
        yield archive


def _fonts(parts: Mapping[str, str | bytes]) -> DocumentFonts:
    with _package(parts) as archive:
        return read_document_fonts(archive)


def _simple_parts(document_body: str) -> dict[str, str | bytes]:
    return {
        "_rels/.rels": _rels(("rId1", "officeDocument", "word/document.xml")),
        "word/document.xml": _document(document_body),
    }


# --- missing_fonts ---


@pytest.mark.req("7.4")
def test_missing_fonts_subtracts_embedded_and_installed_ignoring_case() -> None:
    referenced = {"Calibri", "Arial", "Carlito", "Élite", "Cambria"}
    embedded = {"arial"}
    installed = {"CARLITO", "élite", "DejaVu Sans"}
    assert missing_fonts(referenced, embedded, installed) == ("Calibri", "Cambria")


@pytest.mark.req("7.4")
def test_missing_fonts_reports_each_font_once_in_case_insensitive_order() -> None:
    result = missing_fonts({"calibri", "CALIBRI", "Calibri", "arial", "Bodoni"}, set(), set())
    assert len(result) == 3
    assert [name.casefold() for name in result] == ["arial", "bodoni", "calibri"]
    assert result[2] in {"calibri", "CALIBRI", "Calibri"}
    # No depende del orden en que se recorren los conjuntos
    assert missing_fonts(frozenset({"CALIBRI", "calibri", "Calibri"}), set(), set()) == (
        missing_fonts({"Calibri", "calibri", "CALIBRI"}, set(), set())
    )


@pytest.mark.req("7.4")
def test_missing_fonts_is_empty_when_everything_is_available() -> None:
    assert missing_fonts(set(), {"Arial"}, {"Calibri"}) == ()
    assert missing_fonts({"Arial", "Calibri"}, {"ARIAL"}, {"calibri"}) == ()


# --- read_document_fonts ---


@pytest.mark.req("7.4")
def test_collects_font_table_styles_runs_and_referenced_theme_fonts() -> None:
    body_run = _run('w:ascii="Fuente Run" w:hAnsi="Fuente Run" w:cs="Fuente CS"')
    footnote_run = _run('w:ascii="Fuente Nota"')
    endnote_run = _run('w:eastAsia="Fuente Final"')
    parts: dict[str, str | bytes] = {
        "_rels/.rels": _rels(("rId1", "officeDocument", "word/document.xml")),
        "word/document.xml": _document(body_run),
        "word/_rels/document.xml.rels": _rels(
            ("rId1", "styles", "styles.xml"),
            ("rId2", "fontTable", "fontTable.xml"),
            ("rId3", "theme", "theme/theme1.xml"),
            ("rId4", "header", "header1.xml"),
            ("rId5", "footer", "/word/footer1.xml"),
            ("rId6", "footnotes", "footnotes.xml"),
            ("rId7", "endnotes", "endnotes.xml"),
            ("rId8", "numbering", "numbering.xml"),
        ),
        "word/styles.xml": _wml(
            "styles",
            "<w:docDefaults><w:rPrDefault><w:rPr>"
            '<w:rFonts w:asciiTheme="minorHAnsi" w:hAnsiTheme="minorHAnsi"'
            ' w:eastAsiaTheme="minorEastAsia" w:cstheme="minorBidi"/>'
            "</w:rPr></w:rPrDefault></w:docDefaults>"
            '<w:style w:type="paragraph" w:styleId="Titulo"><w:rPr>'
            '<w:rFonts w:ascii="Fuente Estilo"/></w:rPr></w:style>',
        ),
        "word/fontTable.xml": _wml(
            "fonts",
            '<w:font w:name="Fuente Tabla"><w:family w:val="roman"/></w:font>'
            '<w:font w:name="Fuente Incrustada"><w:embedRegular r:id="rId1"/>'
            '<w:embedBold r:id="rId2"/></w:font>',
        ),
        "word/_rels/fontTable.xml.rels": _rels(
            ("rId1", "font", "fonts/font1.odttf"), ("rId2", "font", "fonts/font2.odttf")
        ),
        "word/fonts/font1.odttf": b"\x00" * 32,
        "word/fonts/font2.odttf": b"\x00" * 32,
        "word/theme/theme1.xml": _theme(
            major="Fuente Mayor", minor="Fuente Menor", minor_ea="Menor EA"
        ),
        "word/header1.xml": _wml("hdr", _run('w:ascii="Fuente Encabezado"')),
        "word/footer1.xml": _wml("ftr", _run('w:hAnsi="Fuente Pie"')),
        "word/footnotes.xml": _wml("footnotes", f"<w:footnote>{footnote_run}</w:footnote>"),
        "word/endnotes.xml": _wml("endnotes", f"<w:endnote>{endnote_run}</w:endnote>"),
        "word/numbering.xml": _wml(
            "numbering",
            '<w:abstractNum w:abstractNumId="0"><w:lvl w:ilvl="0"><w:rPr>'
            '<w:rFonts w:ascii="Symbol" w:hAnsi="Symbol"/></w:rPr></w:lvl></w:abstractNum>',
        ),
        # Partes sin relación con el documento principal: no cuentan
        "word/glossary/document.xml": _document(_run('w:ascii="Fuente Glosario"')),
        "word/header2.xml": _wml("hdr", _run('w:ascii="Fuente Sin Relacion"')),
    }
    fonts = _fonts(parts)
    assert fonts.referenced == {
        "Fuente Run",
        "Fuente CS",
        "Fuente Menor",  # minorHAnsi del tema
        "Menor EA",  # minorEastAsia del tema; minorBidi está vacío y no hay nombre explícito
        "Fuente Estilo",
        "Fuente Tabla",
        "Fuente Incrustada",
        "Fuente Encabezado",
        "Fuente Pie",
        "Fuente Nota",
        "Fuente Final",
        "Symbol",
    }
    assert fonts.embedded == {"Fuente Incrustada"}


@pytest.mark.req("7.4")
def test_theme_reference_replaces_the_explicit_name_unless_the_theme_lacks_it() -> None:
    parts = _simple_parts(
        _run(
            'w:ascii="Arial" w:asciiTheme="minorHAnsi" w:hAnsi="Arial" w:hAnsiTheme="majorHAnsi"'
            ' w:eastAsia="EA Explicita" w:eastAsiaTheme="majorEastAsia"'
            ' w:cs="CS Explicita" w:cstheme="referenciaInexistente"'
        )
    )
    parts["word/_rels/document.xml.rels"] = _rels(("rId1", "theme", "theme/theme1.xml"))
    parts["word/theme/theme1.xml"] = _theme(major="Fuente Mayor", minor="Fuente Menor")
    # Arial no cuenta: las dos posiciones donde aparece tienen una referencia al tema que existe
    assert _fonts(parts).referenced == {
        "Fuente Menor",
        "Fuente Mayor",
        "EA Explicita",  # majorEastAsia está vacío en el tema
        "CS Explicita",  # la referencia no es un valor del tema
    }


@pytest.mark.req("7.4")
def test_theme_references_without_a_theme_use_the_explicit_names() -> None:
    parts = _simple_parts(
        _run('w:ascii="Arial" w:asciiTheme="minorHAnsi" w:hAnsiTheme="minorHAnsi"')
    )
    assert _fonts(parts).referenced == {"Arial"}


@pytest.mark.req("7.4")
def test_embedding_requires_an_internal_relationship_to_an_existing_part() -> None:
    parts = _simple_parts("")
    parts["word/_rels/document.xml.rels"] = _rels(("rId1", "fontTable", "fontTable.xml"))
    parts["word/fontTable.xml"] = _wml(
        "fonts",
        '<w:font w:name="Valida"><w:embedBoldItalic r:id="rId1"/></w:font>'
        '<w:font w:name="Sin Parte"><w:embedItalic r:id="rId2"/></w:font>'
        '<w:font w:name="Externa"><w:embedRegular r:id="rId3"/></w:font>'
        '<w:font w:name="Sin Relacion"><w:embedRegular r:id="rId9"/></w:font>'
        '<w:font w:name="Sin Id"><w:embedRegular/></w:font>',
    )
    parts["word/_rels/fontTable.xml.rels"] = _rels(
        ("rId1", "font", "fonts/valida.odttf"),
        ("rId2", "font", "fonts/falta.odttf"),
        ("rId3", "font", "https://ejemplo.invalid/fuente.odttf"),
        external="rId3",
    )
    parts["word/fonts/valida.odttf"] = b"\x00" * 32
    fonts = _fonts(parts)
    assert fonts.referenced == {"Valida", "Sin Parte", "Externa", "Sin Relacion", "Sin Id"}
    assert fonts.embedded == {"Valida"}


@pytest.mark.req("7.4")
def test_reads_a_python_docx_document() -> None:
    document = docx.Document()
    document.add_paragraph().add_run("hola").font.name = "Comic Sans MS"
    buffer = io.BytesIO()
    document.save(buffer)
    with zipfile.ZipFile(buffer) as archive:
        fonts = read_document_fonts(archive)
    # Tabla de fuentes de la plantilla de python-docx, el tema (Calibri y Cambria) y el run
    assert fonts.referenced == {
        "Symbol",
        "Times New Roman",
        "Cambria",
        "ＭＳ 明朝",
        "Calibri",
        "ＭＳ ゴシック",
        "Courier",
        "Arial",
        "Comic Sans MS",
    }
    assert fonts.embedded == frozenset()


@pytest.mark.req("7.4")
def test_follows_declared_part_names_ignoring_case_in_strict_ooxml() -> None:
    body_run = _run('w:ascii="Fuente Estricta"')
    parts: dict[str, str | bytes] = {
        "_rels/.rels": _rels(("rId1", "officeDocument", "/Doc/Main.xml"), namespace=_R_STRICT),
        "doc/main.xml": _wml("document", f"<w:body>{body_run}</w:body>", w=_W_STRICT, r=_R_STRICT),
        "DOC/_RELS/MAIN.XML.RELS": _rels(("rId1", "fontTable", "fonts.xml"), namespace=_R_STRICT),
        "Doc/Fonts.xml": _wml(
            "fonts",
            '<w:font w:name="Tabla Estricta"><w:embedRegular r:id="rId1"/></w:font>',
            w=_W_STRICT,
            r=_R_STRICT,
        ),
        "doc/_rels/fonts.xml.rels": _rels(("rId1", "font", "f.odttf"), namespace=_R_STRICT),
        "doc/f.odttf": b"\x00" * 32,
        # No es la parte principal declarada
        "word/document.xml": _document(_run('w:ascii="Fuente Ignorada"')),
    }
    fonts = _fonts(parts)
    assert fonts.referenced == {"Fuente Estricta", "Tabla Estricta"}
    assert fonts.embedded == {"Tabla Estricta"}


@pytest.mark.req("7.4")
def test_uses_word_document_xml_when_the_package_declares_no_main_part() -> None:
    parts: dict[str, str | bytes] = {"word/document.xml": _document(_run('w:ascii="Arial"'))}
    assert _fonts(parts).referenced == {"Arial"}


@pytest.mark.req("7.4")
def test_names_are_stripped_and_blank_or_overlong_names_ignored() -> None:
    longest = "F" * MAX_FONT_NAME_LENGTH
    parts = _simple_parts(
        _run(f'w:ascii="  Con Espacios  " w:hAnsi="   " w:eastAsia="{longest}"')
        + _run(f'w:ascii="" w:hAnsi="{longest}G"')
    )
    assert _fonts(parts).referenced == {"Con Espacios", longest}


@pytest.mark.req("7.4")
def test_missing_main_part_is_an_error() -> None:
    parts: dict[str, str | bytes] = {
        "_rels/.rels": _rels(("rId1", "officeDocument", "word/falta.xml")),
        "word/document.xml": _document(""),
    }
    with pytest.raises(FontAnalysisError):
        _fonts(parts)


@pytest.mark.req("7.4")
def test_malformed_xml_is_an_error() -> None:
    parts = _simple_parts("")
    parts["word/document.xml"] = "<w:document><w:body>"
    with pytest.raises(FontAnalysisError):
        _fonts(parts)


@pytest.mark.req("7.4")
def test_external_entities_are_not_resolved(tmp_path: Path) -> None:
    secret = tmp_path / "secreto.txt"
    secret.write_text("Fuente Secreta", encoding="utf-8")
    parts = _simple_parts("")
    parts["word/document.xml"] = (
        f'<!DOCTYPE w:document [<!ENTITY x SYSTEM "{secret.as_uri()}">]>'
        + _document(_run('w:ascii="&x;"'))
    )
    with pytest.raises(FontAnalysisError):
        _fonts(parts)


# --- fontconfig ---


@pytest.mark.req("7.4")
def test_fc_match_pattern_escapes_the_family_delimiters() -> None:
    assert fc_match_pattern("Calibri Light") == "Calibri Light"
    assert fc_match_pattern("-a:b,c\\d") == "\\-a\\:b\\,c\\\\d"


@pytest.mark.req("7.4")
def test_parse_installed_families_reads_one_family_per_line() -> None:
    output = "DejaVu Sans\nDejaVu Sans Condensed\n\nDejaVu Sans\n  Carlito \n"
    assert parse_installed_families(output) == {"DejaVu Sans", "DejaVu Sans Condensed", "Carlito"}
    assert parse_installed_families("") == frozenset()


@pytest.mark.req("7.4")
@pytest.mark.requires_tool("fc-list", "fc-match")
def test_missing_fonts_get_an_installed_replacement() -> None:
    installed = installed_families()
    assert installed
    # Una familia instalada, escrita con otras mayúsculas, no se informa como reemplazada
    present = html.escape(sorted(installed)[0].upper(), quote=True)
    parts = _simple_parts(_run(f'w:ascii="{present}" w:hAnsi="Fuente Inexistente-12:bold,Otra"'))
    with _package(parts) as archive:
        replacements = find_replacements(archive)
    assert [item.name for item in replacements] == ["Fuente Inexistente-12:bold,Otra"]
    assert replacements[0].replacement in installed


@pytest.mark.req("7.4")
@pytest.mark.requires_tool("fc-list", "fc-match")
def test_calibri_is_replaced_by_its_metric_compatible_font() -> None:
    if "Carlito" not in installed_families():
        pytest.skip("Carlito no está instalada")
    assert match_font("Calibri") == "Carlito"
