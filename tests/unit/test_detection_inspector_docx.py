"""Pruebas del inspector `docx` en src/converter/detection/inspectors/docx.py (Req 7.5, 7.6, 7.7,
16.1, 16.2).

Los paquetes se arman con `zipfile` y con las partes de WordprocessingML que importan en cada
caso; algunos se generan con python-docx. No se usan archivos externos. Los Límites_de_Descompresión
son los reales (fijos, D18): los paquetes que los superan son pequeños porque se escriben con
datos muy compresibles, con una cantidad excesiva de entradas o con tamaños declarados falsos.
"""

from __future__ import annotations

import io
import socket
import struct
import subprocess
import sys
import zipfile
from collections.abc import Iterable, Mapping, Sequence
from pathlib import Path

import docx
import pytest

from converter.detection.inspectors import InspectorInput, get_inspector, registered_inspectors
from converter.detection.inspectors.docx import DocxInfo, inspect_docx
from converter.detection.source_info import (
    InspectorOutputError,
    SourceInfo,
    SourceRejection,
    dump_inspector_output,
    load_inspector_output,
    registered_source_infos,
)
from converter.detection.zipcheck import ExceededLimit, check_zip_package
from converter.messages_es import FailureReason
from converter.models import Stage
from converter.sandbox.limits import SandboxTool
from converter.workers.inspect import MODULE
from tests.format_samples import (
    WORDPROCESSING_MAIN,
    content_types,
    encrypted_docx,
    zip_package,
)

_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
_W_STRICT = "http://purl.oclc.org/ooxml/wordprocessingml/main"
_R_STRICT = "http://purl.oclc.org/ooxml/officeDocument/relationships"
_PKG = "http://schemas.openxmlformats.org/package/2006/relationships"
_VBA_PROJECT = "http://schemas.microsoft.com/office/2006/relationships/vbaProject"

NO_CONTENT = DocxInfo(
    has_macros=False, has_comments=False, has_tracked_changes=False, external_relationships=0
)

type Rel = tuple[str, str, bool]  # tipo completo, destino y si es externa


def rel(kind: str, target: str, *, external: bool = False, namespace: str = _R) -> Rel:
    """Relación de Office de tipo `kind` (sin prefijo)."""
    return (f"{namespace}/{kind}", target, external)


def rels_part(*relationships: Rel) -> bytes:
    items = "".join(
        f'<Relationship Id="rId{n}" Type="{kind}" Target="{target}"'
        + (' TargetMode="External"' if external else "")
        + "/>"
        for n, (kind, target, external) in enumerate(relationships, 1)
    )
    return (
        f'<?xml version="1.0" encoding="UTF-8"?><Relationships xmlns="{_PKG}">{items}'
        "</Relationships>"
    ).encode()


def wml(root: str, body: str, *, w: str = _W) -> bytes:
    """Parte de WordprocessingML con la raíz `w:root` y `body` como contenido."""
    return f'<?xml version="1.0"?><w:{root} xmlns:w="{w}">{body}</w:{root}>'.encode()


def package(
    body: str = "<w:p><w:r><w:t>hola</w:t></w:r></w:p>",
    *,
    w: str = _W,
    relationships: Sequence[Rel] = (),
    parts: Mapping[str, bytes] | None = None,
    package_relationships: Sequence[Rel] | None = None,
) -> bytes:
    """DOCX con `body` en la parte principal; `parts` agrega partes y `relationships`, las
    relaciones de la parte principal. Sin `package_relationships`, el paquete declara
    `word/document.xml` como parte principal."""
    if package_relationships is None:
        package_relationships = [rel("officeDocument", "word/document.xml")]
    return zip_package(
        ("[Content_Types].xml", content_types(WORDPROCESSING_MAIN)),
        ("_rels/.rels", rels_part(*package_relationships)),
        ("word/document.xml", wml("document", f"<w:body>{body}</w:body>", w=w)),
        ("word/_rels/document.xml.rels", rels_part(*relationships)),
        *(parts or {}).items(),
    )


def inspect_bytes(tmp_path: Path, data: bytes) -> SourceInfo | SourceRejection:
    path = tmp_path / "source"
    path.write_bytes(data)
    return inspect_docx(InspectorInput(path=path, real_type="DOCX", max_image_pixels=1_000_000))


def inspect_info(tmp_path: Path, data: bytes) -> DocxInfo:
    result = inspect_bytes(tmp_path, data)
    assert isinstance(result, DocxInfo), result
    return result


def assert_rejected(result: SourceInfo | SourceRejection, reason: FailureReason) -> None:
    assert isinstance(result, SourceRejection), result
    assert result.reason is reason
    assert result.failure.stage is Stage.DETECTION
    assert dict(result.params) == {}


# --- Registro ---


def test_inspector_is_registered_with_the_catalog_name() -> None:
    assert "docx" in registered_inspectors()
    spec = get_inspector("docx")
    assert spec.function is inspect_docx
    assert spec.info_type is DocxInfo
    # El lector ZIP y lxml no tienen una herramienta con salida que sanear
    assert spec.tool is None
    assert spec.sandbox_tool is SandboxTool.PYTHON
    assert registered_source_infos()["docx"] is DocxInfo


def test_importing_the_inspectors_does_not_load_lxml() -> None:
    # El proceso principal y cada worker importan todos los inspectores solo para conocer sus
    # nombres; lxml se importa al inspeccionar un DOCX
    code = (
        "import sys, converter.detection.inspectors, converter.workers.inspect;"
        "sys.exit(int('lxml' in sys.modules))"
    )

    completed = subprocess.run(
        [sys.executable, "-I", "-c", code],
        capture_output=True,
        text=True,
        timeout=60,
        check=False,
    )

    assert completed.returncode == 0, completed.stderr


# --- Contenido del paquete ---


def test_package_without_macros_comments_changes_or_external_relationships(
    tmp_path: Path,
) -> None:
    assert inspect_info(tmp_path, package()) == NO_CONTENT


def test_document_written_by_python_docx_has_nothing_to_omit(tmp_path: Path) -> None:
    document = docx.Document()
    document.add_heading("Informe", level=1)
    document.add_paragraph("Texto del informe con ñ y acentos: canción.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "A"
    table.cell(1, 1).text = "B"
    buffer = io.BytesIO()
    document.save(buffer)

    assert inspect_info(tmp_path, buffer.getvalue()) == NO_CONTENT


@pytest.mark.req("7.6")
def test_document_with_a_comment_written_by_python_docx_has_comments(tmp_path: Path) -> None:
    document = docx.Document()
    paragraph = document.add_paragraph("Texto revisado por una persona.")
    document.add_comment(paragraph.runs, text="Revisar este párrafo.", author="Ana", initials="A")
    buffer = io.BytesIO()
    document.save(buffer)

    info = inspect_info(tmp_path, buffer.getvalue())

    assert info == DocxInfo(
        has_macros=False, has_comments=True, has_tracked_changes=False, external_relationships=0
    )


# Control de cambios


def tracked(name: str) -> str:
    return f'<w:{name} w:id="1" w:author="Ana" w:date="2025-01-01T00:00:00Z"/>'


TRACKED_CHANGE_NAMES = [
    "ins",
    "del",
    "moveFrom",
    "moveTo",
    "rPrChange",
    "pPrChange",
    "sectPrChange",
    "tblPrChange",
    "tblPrExChange",
    "tblGridChange",
    "trPrChange",
    "tcPrChange",
    "numberingChange",
    "cellIns",
    "cellDel",
    "cellMerge",
]


@pytest.mark.req("7.6")
@pytest.mark.parametrize("name", TRACKED_CHANGE_NAMES)
@pytest.mark.parametrize("namespace", [_W, _W_STRICT], ids=["transicion", "estricto"])
def test_pending_changes_in_the_main_part_are_detected(
    tmp_path: Path, name: str, namespace: str
) -> None:
    info = inspect_info(tmp_path, package(f"<w:p>{tracked(name)}</w:p>", w=namespace))

    assert info.has_tracked_changes is True
    assert info == DocxInfo(
        has_macros=False, has_comments=False, has_tracked_changes=True, external_relationships=0
    )


@pytest.mark.req("7.6")
def test_insertions_and_deletions_with_text_are_detected(tmp_path: Path) -> None:
    body = (
        "<w:p><w:r><w:t>Texto </w:t></w:r>"
        '<w:ins w:id="1" w:author="Ana"><w:r><w:t>nuevo</w:t></w:r></w:ins>'
        '<w:del w:id="2" w:author="Ana"><w:r><w:delText>viejo</w:delText></w:r></w:del></w:p>'
    )

    assert inspect_info(tmp_path, package(body)).has_tracked_changes is True


@pytest.mark.req("7.6")
@pytest.mark.parametrize("kind", ["header", "footer", "footnotes", "endnotes"])
def test_pending_changes_in_a_related_part_are_detected(tmp_path: Path, kind: str) -> None:
    part = f"word/{kind}1.xml"
    data = package(
        relationships=[rel(kind, f"{kind}1.xml")],
        parts={part: wml(kind, f"<w:p>{tracked('ins')}</w:p>")},
    )

    assert inspect_info(tmp_path, data).has_tracked_changes is True


@pytest.mark.parametrize(
    "case",
    [
        pytest.param("parte_sin_relacion", id="parte_sin_relacion"),
        pytest.param("estilos", id="estilos"),
        pytest.param("otro_espacio_de_nombres", id="otro_espacio_de_nombres"),
        pytest.param("control_activado", id="control_activado"),
        pytest.param("texto_eliminado_sin_marca", id="texto_eliminado_sin_marca"),
    ],
)
def test_content_without_pending_changes_is_not_reported(tmp_path: Path, case: str) -> None:
    changed = wml("header", f"<w:p>{tracked('ins')}</w:p>")
    if case == "parte_sin_relacion":
        # El encabezado está en el paquete, pero ninguna relación de la parte principal lo usa
        data = package(parts={"word/header1.xml": changed})
    elif case == "estilos":
        # Los estilos no se leen: un `w:rPrChange` ahí no es un cambio del contenido
        data = package(
            relationships=[rel("styles", "styles.xml")],
            parts={"word/styles.xml": wml("styles", tracked("rPrChange"))},
        )
    elif case == "otro_espacio_de_nombres":
        data = package('<p xmlns="urn:otro"><ins id="1"/></p>')
    elif case == "control_activado":
        # `w:trackRevisions` solo activa el control de cambios para lo que se edite después
        data = package(
            relationships=[rel("settings", "settings.xml")],
            parts={"word/settings.xml": wml("settings", "<w:trackRevisions/>")},
        )
    else:
        data = package("<w:p><w:r><w:delText>texto</w:delText></w:r></w:p>")

    assert inspect_info(tmp_path, data) == NO_CONTENT


# Comentarios


@pytest.mark.req("7.6")
@pytest.mark.parametrize("namespace", [_W, _W_STRICT], ids=["transicion", "estricto"])
def test_comments_part_with_a_comment_is_detected(tmp_path: Path, namespace: str) -> None:
    comment = (
        '<w:comment w:id="0" w:author="Ana"><w:p><w:r><w:t>Revisar</w:t></w:r></w:p></w:comment>'
    )
    data = package(
        relationships=[rel("comments", "comments.xml")],
        parts={"word/comments.xml": wml("comments", comment, w=namespace)},
    )

    assert inspect_info(tmp_path, data).has_comments is True


@pytest.mark.parametrize(
    "case", ["parte_vacia", "referencia_sin_parte", "parte_sin_relacion", "destino_inexistente"]
)
def test_document_without_a_comment_is_not_reported(tmp_path: Path, case: str) -> None:
    comment = wml("comments", '<w:comment w:id="0" w:author="Ana"/>')
    if case == "parte_vacia":
        # Word deja una parte de comentarios vacía al borrar el último comentario
        data = package(
            relationships=[rel("comments", "comments.xml")],
            parts={"word/comments.xml": wml("comments", "")},
        )
    elif case == "referencia_sin_parte":
        body = '<w:p><w:commentRangeStart w:id="0"/><w:commentRangeEnd w:id="0"/></w:p>'
        data = package(body)
    elif case == "parte_sin_relacion":
        data = package(parts={"word/comments.xml": comment})
    else:
        data = package(relationships=[rel("comments", "no_existe.xml")])

    assert inspect_info(tmp_path, data).has_comments is False


# Macros


@pytest.mark.req("7.5")
@pytest.mark.parametrize(
    ("relationships", "parts"),
    [
        pytest.param(
            [("vbaProject", "vbaProject.bin")],
            {"word/vbaProject.bin": b"\xd0\xcf\x11\xe0"},
            id="parte_y_relacion",
        ),
        pytest.param([], {"word/vbaProject.bin": b"\xd0\xcf\x11\xe0"}, id="solo_la_parte"),
        pytest.param([("vbaProject", "vbaProject.bin")], {}, id="solo_la_relacion"),
        pytest.param([], {"word/VBAPROJECT.BIN": b"\xd0\xcf\x11\xe0"}, id="mayusculas"),
    ],
)
def test_vba_project_is_detected_as_macros(
    tmp_path: Path, relationships: list[tuple[str, str]], parts: dict[str, bytes]
) -> None:
    rels = [
        (_VBA_PROJECT, target, False) if kind == "vbaProject" else rel(kind, target)
        for kind, target in relationships
    ]

    info = inspect_info(tmp_path, package(relationships=rels, parts=parts))

    assert info == DocxInfo(
        has_macros=True, has_comments=False, has_tracked_changes=False, external_relationships=0
    )


@pytest.mark.req("7.5")
def test_activex_controls_are_not_macros(tmp_path: Path) -> None:
    data = package(
        relationships=[rel("control", "activeX/activeX1.xml")],
        parts={
            "word/activeX/activeX1.xml": wml("ocx", ""),
            "word/activeX/activeX1.bin": b"\x00" * 16,
        },
    )

    assert inspect_info(tmp_path, data).has_macros is False


# Relaciones externas


@pytest.mark.req("7.5")
def test_external_relationships_that_are_not_hyperlinks_are_counted_in_every_part(
    tmp_path: Path,
) -> None:
    data = package(
        relationships=[
            rel("image", "file:///C:/imagenes/logo.png", external=True),
            rel("oleObject", "\\\\servidor\\recurso\\hoja.xlsx", external=True),
            rel("hyperlink", "https://example.com/", external=True),
            rel("hyperlink", "mailto:ana@example.com", external=True),
            rel("image", "media/image1.png"),  # interna: está dentro del paquete
            rel("settings", "settings.xml"),
            rel("header", "header1.xml"),
        ],
        parts={
            "word/settings.xml": wml("settings", ""),
            "word/_rels/settings.xml.rels": rels_part(
                rel("attachedTemplate", "https://example.com/plantilla.dotm", external=True)
            ),
            "word/header1.xml": wml("header", "<w:p/>"),
            "word/_rels/header1.xml.rels": rels_part(
                rel("image", "http://example.com/membrete.png", external=True),
                rel("hyperlink", "https://example.com/otro", external=True),
            ),
        },
    )

    # Imagen y objeto de la parte principal, plantilla de la configuración e imagen del encabezado
    assert inspect_info(tmp_path, data).external_relationships == 4


@pytest.mark.req("7.5")
def test_external_relationship_of_the_package_is_counted(tmp_path: Path) -> None:
    data = package(
        package_relationships=[
            rel("officeDocument", "word/document.xml"),
            rel("thumbnail", "http://example.com/miniatura.png", external=True),
        ]
    )

    assert inspect_info(tmp_path, data).external_relationships == 1


@pytest.mark.req("7.5")
def test_strict_relation_types_are_recognized(tmp_path: Path) -> None:
    data = package(
        w=_W_STRICT,
        package_relationships=[rel("officeDocument", "word/document.xml", namespace=_R_STRICT)],
        relationships=[
            rel("hyperlink", "https://example.com/", external=True, namespace=_R_STRICT),
            rel("image", "https://example.com/logo.png", external=True, namespace=_R_STRICT),
        ],
    )

    assert inspect_info(tmp_path, data).external_relationships == 1


def test_relationship_marked_as_internal_is_not_external(tmp_path: Path) -> None:
    rels = (
        f'<?xml version="1.0"?><Relationships xmlns="{_PKG}">'
        f'<Relationship Id="rId1" Type="{_R}/image" Target="media/image1.png" '
        'TargetMode="Internal"/>'
        f'<Relationship Id="rId2" Type="{_R}/image" Target="media/image2.png"/>'
        "</Relationships>"
    ).encode()
    data = package(parts={"word/media/image1.png": b"\x89PNG", "word/media/image2.png": b"\x89PNG"})
    data = replace_part(data, "word/_rels/document.xml.rels", rels)

    assert inspect_info(tmp_path, data).external_relationships == 0


def test_main_part_is_the_one_the_package_declares(tmp_path: Path) -> None:
    # Los cambios están en `word/principal.xml`, no en `word/document.xml`
    data = zip_package(
        ("[Content_Types].xml", content_types(WORDPROCESSING_MAIN)),
        ("_rels/.rels", rels_part(rel("officeDocument", "/word/principal.xml"))),
        ("word/document.xml", wml("document", "<w:body/>")),
        ("word/principal.xml", wml("document", f"<w:body><w:p>{tracked('ins')}</w:p></w:body>")),
        ("word/_rels/principal.xml.rels", rels_part(rel("comments", "comentarios.xml"))),
        (
            "word/comentarios.xml",
            wml("comments", '<w:comment w:id="0" w:author="Ana"/>'),
        ),
    )

    info = inspect_info(tmp_path, data)

    assert info.has_tracked_changes is True
    assert info.has_comments is True


def test_package_without_relationships_uses_word_document_xml(tmp_path: Path) -> None:
    data = zip_package(
        ("[Content_Types].xml", content_types(WORDPROCESSING_MAIN)),
        ("word/document.xml", wml("document", f"<w:body><w:p>{tracked('del')}</w:p></w:body>")),
    )

    assert inspect_info(tmp_path, data).has_tracked_changes is True


# --- Límites_de_Descompresión ---


def zeros_entry(size: int = 200_000) -> tuple[str, bytes]:
    """Entrada muy compresible: con deflate su Razón_de_Compresión supera 200."""
    return ("word/media/vacio.bin", bytes(size))


def with_entries(extra: Iterable[tuple[str, bytes]]) -> bytes:
    return zip_package(
        ("[Content_Types].xml", content_types(WORDPROCESSING_MAIN)),
        ("_rels/.rels", rels_part(rel("officeDocument", "word/document.xml"))),
        ("word/document.xml", wml("document", "<w:body/>")),
        *extra,
    )


def entry_count(data: bytes) -> int:
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        return len(archive.namelist())


def patch_declared_size(data: bytes, name: str, size: int) -> bytes:
    """Cambia el tamaño descomprimido que el directorio central declara para la entrada `name`."""
    encoded = name.encode()
    position = data.find(b"PK\x01\x02")
    while position >= 0:
        (name_length,) = struct.unpack_from("<H", data, position + 28)
        if data[position + 46 : position + 46 + name_length] == encoded:
            patched = bytearray(data)
            struct.pack_into("<L", patched, position + 24, size)
            return bytes(patched)
        position = data.find(b"PK\x01\x02", position + 4)
    raise AssertionError(f"el directorio central no tiene la entrada {name}")


def replace_part(data: bytes, name: str, content: bytes) -> bytes:
    """Copia de `data` con la parte `name` reemplazada por `content`."""
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        entries = [(info.filename, archive.read(info)) for info in archive.infolist()]
    return zip_package(*[(n, content if n == name else d) for n, d in entries])


@pytest.mark.req("7.7", "16.1")
def test_entry_with_a_compression_ratio_above_the_limit_is_rejected(tmp_path: Path) -> None:
    data = with_entries([zeros_entry()])

    result = inspect_bytes(tmp_path, data)

    assert_rejected(result, FailureReason.DECOMPRESSION_LIMITS)


@pytest.mark.req("7.7", "16.1")
def test_more_than_ten_thousand_entries_are_rejected(tmp_path: Path) -> None:
    fillers = [(f"word/relleno/{n}.txt", b"") for n in range(10_001 - 3)]
    data = with_entries(fillers)

    assert entry_count(data) == 10_001
    assert_rejected(inspect_bytes(tmp_path, data), FailureReason.DECOMPRESSION_LIMITS)


@pytest.mark.req("7.7")
def test_exactly_ten_thousand_entries_are_accepted(tmp_path: Path) -> None:
    fillers = [(f"word/relleno/{n}.txt", b"") for n in range(10_000 - 3)]
    data = with_entries(fillers)

    assert entry_count(data) == 10_000
    assert inspect_info(tmp_path, data) == NO_CONTENT


@pytest.mark.req("7.7", "16.1")
def test_declared_sizes_above_the_total_limit_are_rejected(tmp_path: Path) -> None:
    data = patch_declared_size(with_entries([]), "word/document.xml", 512_000_001)

    assert_rejected(inspect_bytes(tmp_path, data), FailureReason.DECOMPRESSION_LIMITS)


@pytest.mark.req("7.7", "16.2")
def test_entry_that_produces_more_than_it_declares_is_rejected_by_the_limits(
    tmp_path: Path,
) -> None:
    # El directorio central declara 1 000 bytes y la entrada produce 200 000
    name, content = zeros_entry()
    data = patch_declared_size(with_entries([(name, content)]), name, 1_000)

    assert_rejected(inspect_bytes(tmp_path, data), FailureReason.DECOMPRESSION_LIMITS)


def damaged_and_over_the_limits() -> bytes:
    """Paquete con una entrada dañada (CRC) y otra que supera la razón de compresión."""
    buffer = io.BytesIO()
    with zipfile.ZipFile(buffer, "w") as archive:
        archive.writestr("word/nota.txt", b"texto de la nota " * 4, zipfile.ZIP_STORED)
        name, content = zeros_entry()
        archive.writestr(name, content, zipfile.ZIP_DEFLATED)
    data = buffer.getvalue().replace(b"texto de la nota", b"Texto de la nota", 1)
    return patch_declared_size(data, "word/media/vacio.bin", 1_000)


@pytest.mark.req("7.7", "16.1", "16.2")
def test_exceeding_the_limits_has_priority_over_damage(tmp_path: Path) -> None:
    data = damaged_and_over_the_limits()
    verdict = check_zip_package(io.BytesIO(data))
    # El paquete cumple las dos condiciones: es la prioridad la que decide el rechazo
    assert verdict.exceeded is ExceededLimit.COMPRESSION_RATIO
    assert verdict.damaged is True

    result = inspect_bytes(tmp_path, data)

    assert_rejected(result, FailureReason.DECOMPRESSION_LIMITS)


# --- Archivo dañado ---


def stored_package_with_crc_error() -> bytes:
    data = zip_package(
        ("[Content_Types].xml", content_types(WORDPROCESSING_MAIN)),
        ("_rels/.rels", rels_part(rel("officeDocument", "word/document.xml"))),
        ("word/document.xml", wml("document", "<w:body><w:p>texto</w:p></w:body>")),
        compression=zipfile.ZIP_STORED,
    )
    return data.replace(b"<w:p>texto", b"<w:p>Texto", 1)


def package_with_invalid_utf8_name() -> bytes:
    # zipfile escribe "ñ" como UTF-8 con el indicador 0x800; los bytes C3 B1 pasan a ser inválidos
    data = package(parts={"word/ñ.bin": b"datos"})
    assert data.count("ñ".encode()) == 2  # cabecera local y directorio central
    return data.replace("ñ".encode(), b"\xff\xfe")


def damaged_samples() -> dict[str, bytes]:
    valid = package()
    return {
        "vacio": b"",
        "no_es_un_zip": b"esto no es un paquete ZIP",
        "truncado": valid[:-10],
        "primera_mitad": valid[: len(valid) // 2],
        "crc_incorrecto": stored_package_with_crc_error(),
        "nombre_con_utf8_invalido": package_with_invalid_utf8_name(),
        "metodo_no_admitido": zip_package(
            ("[Content_Types].xml", content_types(WORDPROCESSING_MAIN)),
            ("word/document.xml", wml("document", "<w:body/>")),
            compression=zipfile.ZIP_LZMA,
        ),
        "paquete_vacio": zip_package(),
        "sin_parte_principal": zip_package(
            ("[Content_Types].xml", content_types(WORDPROCESSING_MAIN)),
            ("_rels/.rels", rels_part(rel("officeDocument", "word/document.xml"))),
        ),
        "parte_principal_declarada_inexistente": package(
            package_relationships=[rel("officeDocument", "word/principal.xml")]
        ),
        "parte_principal_vacia": replace_part(package(), "word/document.xml", b""),
        "parte_principal_mal_formada": replace_part(
            package(), "word/document.xml", b"<w:document><w:body></w:document>"
        ),
        "parte_principal_sin_cierre": replace_part(
            package(), "word/document.xml", wml("document", "<w:body>")[: -len(b"</w:document>")]
        ),
        "relaciones_del_paquete_mal_formadas": replace_part(
            package(), "_rels/.rels", b"<Relationships><Relationship>"
        ),
        "relaciones_de_la_principal_mal_formadas": replace_part(
            package(), "word/_rels/document.xml.rels", b"\x00\x01\x02"
        ),
        "relaciones_de_otra_parte_mal_formadas": package(
            parts={"word/_rels/settings.xml.rels": b"no es XML"}
        ),
        "encabezado_relacionado_mal_formado": package(
            relationships=[rel("header", "header1.xml")],
            parts={"word/header1.xml": b"<w:hdr xmlns:w='urn:x'><w:p></w:hdr>"},
        ),
        "comentarios_mal_formados": package(
            relationships=[rel("comments", "comments.xml")],
            parts={"word/comments.xml": b"<w:comments"},
        ),
        "codificacion_desconocida": replace_part(
            package(), "word/document.xml", b'<?xml version="1.0" encoding="no-existe"?><a/>'
        ),
        "contenido_despues_de_la_raiz": replace_part(
            package(), "word/document.xml", wml("document", "") + b"<otra/>"
        ),
    }


@pytest.mark.req("16.1", "16.2")
@pytest.mark.parametrize("name", list(damaged_samples()))
def test_package_with_an_unreadable_structure_is_rejected_as_damaged(
    tmp_path: Path, name: str
) -> None:
    result = inspect_bytes(tmp_path, damaged_samples()[name])

    assert_rejected(result, FailureReason.DAMAGED)


def test_encrypted_docx_container_is_not_a_zip_package(tmp_path: Path) -> None:
    # Un DOCX_Cifrado lo rechaza la sonda de variante antes de llegar aquí; si llegara, un archivo
    # OLE no es un paquete ZIP
    assert_rejected(inspect_bytes(tmp_path, encrypted_docx()), FailureReason.DAMAGED)


def test_part_that_the_inspector_does_not_read_is_not_validated(tmp_path: Path) -> None:
    # La conversión detecta el daño de las partes que no se leen (Req 16.2, "desde en_proceso")
    data = package(
        relationships=[rel("styles", "styles.xml")], parts={"word/styles.xml": b"<w:styles>"}
    )

    assert inspect_info(tmp_path, data) == NO_CONTENT


# --- Seguridad del XML ---


@pytest.mark.req("16.2")
def test_external_entity_is_not_read(tmp_path: Path) -> None:
    # El archivo al que apunta la entidad no es XML bien formado: si el analizador lo leyera,
    # la parte estaría dañada
    secret = tmp_path / "secreto.txt"
    secret.write_text("<<< no es XML")
    body = f'<!DOCTYPE d [<!ENTITY x SYSTEM "{secret.as_uri()}">]>'
    document = (
        f'<?xml version="1.0"?>{body}<w:document xmlns:w="{_W}"><w:body>&x;</w:body></w:document>'
    )

    data = replace_part(package(), "word/document.xml", document.encode())

    assert inspect_info(tmp_path, data) == NO_CONTENT


@pytest.mark.req("16.2")
def test_external_entity_does_not_open_a_network_connection(tmp_path: Path) -> None:
    # El libxml2 de las ruedas de lxml no incluye cliente HTTP, así que esta prueba no distingue
    # las opciones del analizador; protege contra una compilación de lxml que sí lo incluya
    with socket.socket() as server:
        server.bind(("127.0.0.1", 0))
        server.listen()
        server.setblocking(False)
        port = server.getsockname()[1]
        body = f'<!DOCTYPE d [<!ENTITY x SYSTEM "http://127.0.0.1:{port}/x">]>'
        document = (
            f'<?xml version="1.0"?>{body}<w:document xmlns:w="{_W}">'
            "<w:body>&x;</w:body></w:document>"
        )
        data = replace_part(package(), "word/document.xml", document.encode())

        info = inspect_info(tmp_path, data)

        assert info == NO_CONTENT
        with pytest.raises(BlockingIOError):
            server.accept()


@pytest.mark.req("16.2")
def test_entity_expansion_bomb_does_not_expand(tmp_path: Path) -> None:
    entities = '<!ENTITY e0 "lol">' + "".join(
        f'<!ENTITY e{n} "{f"&e{n - 1};" * 10}">' for n in range(1, 10)
    )
    document = (
        f'<?xml version="1.0"?><!DOCTYPE d [{entities}]>'
        f'<w:document xmlns:w="{_W}"><w:body>&e9;</w:body></w:document>'
    )
    data = replace_part(package(), "word/document.xml", document.encode())

    result = inspect_bytes(tmp_path, data)

    # Sin expansión no hay nada que leer; libxml2 puede además dar el documento por dañado
    assert result == NO_CONTENT or result == SourceRejection(FailureReason.DAMAGED)


# --- Datos del SourceInfo ---


def test_info_rejects_values_of_other_types() -> None:
    with pytest.raises(TypeError, match="has_macros"):
        DocxInfo(
            has_macros=1,  # type: ignore[arg-type]
            has_comments=False,
            has_tracked_changes=False,
            external_relationships=0,
        )
    with pytest.raises(TypeError, match="has_tracked_changes"):
        DocxInfo(
            has_macros=False,
            has_comments=False,
            has_tracked_changes="si",  # type: ignore[arg-type]
            external_relationships=0,
        )
    with pytest.raises(TypeError, match="entero"):
        DocxInfo(
            has_macros=False,
            has_comments=False,
            has_tracked_changes=False,
            external_relationships=True,
        )
    with pytest.raises(TypeError, match="entero"):
        DocxInfo(
            has_macros=False,
            has_comments=False,
            has_tracked_changes=False,
            external_relationships="2",  # type: ignore[arg-type]
        )
    with pytest.raises(ValueError, match="negativo"):
        DocxInfo(
            has_macros=False,
            has_comments=False,
            has_tracked_changes=False,
            external_relationships=-1,
        )


@pytest.mark.parametrize(
    "info",
    [
        NO_CONTENT,
        DocxInfo(
            has_macros=True, has_comments=True, has_tracked_changes=True, external_relationships=7
        ),
    ],
)
def test_info_round_trips_through_the_json_of_the_worker(info: DocxInfo) -> None:
    assert load_inspector_output(dump_inspector_output(info)) == info


@pytest.mark.parametrize(
    "data",
    [
        '{"has_macros":false,"has_comments":false,"has_tracked_changes":false}',
        '{"has_macros":0,"has_comments":false,"has_tracked_changes":false,'
        '"external_relationships":0}',
        '{"has_macros":false,"has_comments":false,"has_tracked_changes":false,'
        '"external_relationships":-1}',
        '{"has_macros":false,"has_comments":false,"has_tracked_changes":false,'
        '"external_relationships":true}',
        '{"has_macros":false,"has_comments":false,"has_tracked_changes":false,'
        '"external_relationships":0,"otro":1}',
    ],
)
def test_json_with_invalid_fields_is_rejected(data: str) -> None:
    with pytest.raises(InspectorOutputError):
        load_inspector_output(f'{{"info":{{"type":"docx","data":{data}}}}}')


# --- Worker ---


def run_worker(path: Path) -> SourceInfo | SourceRejection:
    completed = subprocess.run(
        [sys.executable, "-I", "-m", MODULE, "--real-type", "DOCX", "--", "docx", str(path)],
        capture_output=True,
        text=True,
        timeout=60,
        check=False,
    )
    assert completed.returncode == 0, completed.stderr
    return load_inspector_output(completed.stdout)


@pytest.mark.req("7.5", "7.6", "7.7", "16.2")
def test_worker_writes_the_result_of_the_docx_inspector(tmp_path: Path) -> None:
    complete = tmp_path / "completo.docx"
    complete.write_bytes(
        package(
            f"<w:p>{tracked('ins')}</w:p>",
            relationships=[
                rel("comments", "comments.xml"),
                rel("image", "http://example.com/logo.png", external=True),
                (_VBA_PROJECT, "vbaProject.bin", False),
            ],
            parts={
                "word/comments.xml": wml("comments", '<w:comment w:id="0" w:author="Ana"/>'),
                "word/vbaProject.bin": b"\xd0\xcf\x11\xe0",
            },
        )
    )
    over_the_limits = tmp_path / "limites.docx"
    over_the_limits.write_bytes(with_entries([zeros_entry()]))
    damaged = tmp_path / "danado.docx"
    damaged.write_bytes(package()[:-10])

    assert run_worker(complete) == DocxInfo(
        has_macros=True, has_comments=True, has_tracked_changes=True, external_relationships=1
    )
    assert run_worker(over_the_limits) == SourceRejection(FailureReason.DECOMPRESSION_LIMITS)
    assert run_worker(damaged) == SourceRejection(FailureReason.DAMAGED)
